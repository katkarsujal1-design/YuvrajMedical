"""Read-only live checks: OWNER_LIVE_TESTS=1 python -m unittest discover -s tests -p test_owner_reporting.py -v."""
import os
import unittest
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from services.owner_dashboard import load_dashboard, revenue_data, now_ist, display_date, money, rows
from services.owner_reports import build_report, csv_bytes, excel_bytes, pdf_bytes, date_filters, spreadsheet_cell


class ReportUnitTests(unittest.TestCase):
    def test_invalid_dates(self):
        for filters in [{'from':'2026-01-01'}, {'from':'bad','to':'2026-01-01'},
                        {'from':'2026-09-08','to':'2026-09-07'}, {'from':'2020-01-01','to':'2026-01-01'}]:
            with self.subTest(filters=filters), self.assertRaises(ValueError): date_filters(filters)

    def test_currency_null_and_timezone(self):
        self.assertEqual(money(None), '—')
        self.assertEqual(money(Decimal('1250')), '₹1,250.00')
        self.assertEqual(display_date(datetime(2026,9,6,20,0)), '07 Sep 2026, 01:30 AM')
        self.assertEqual(display_date(date(2026,9,7)), '07 Sep 2026')

    def test_spreadsheet_injection(self):
        for value in ['=1+1', '+cmd', '-cmd', '@SUM(A1)', '  =HYPERLINK("x")']:
            self.assertTrue(spreadsheet_cell(value).startswith("'"))
        self.assertEqual(spreadsheet_cell(Decimal('-12')), Decimal('-12'))


@unittest.skipUnless(os.getenv('OWNER_LIVE_TESTS') == '1', 'Explicit live read-only database opt-in required')
class OwnerLiveTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import mysql.connector
        cls.db=mysql.connector.connect(host=os.environ['DB_HOST'],user=os.environ['DB_USER'],
            password=os.environ['DB_PASSWORD'],database=os.environ['DB_NAME'])
        cls.db.start_transaction(readonly=True, consistent_snapshot=True)
        cls.data=load_dashboard(cls.db,{})

    @classmethod
    def tearDownClass(cls):
        cls.db.rollback(); cls.db.close()

    def test_every_metric_against_rows(self):
        metrics=self.data['metrics']; today=now_ist().date()
        order_rows=rows(self.db,'SELECT id,status,payment_status,refund_status,date FROM orders')
        payments=rows(self.db,'SELECT id,order_id,status,amount FROM payments ORDER BY id')
        latest={p['order_id']:p for p in payments}
        eligible=[o for o in order_rows if o['payment_status']=='Verified' and o['status'] not in ('Cancelled','Refunded')
                  and o['refund_status'] not in ('Processing','Completed') and latest.get(o['id'],{}).get('status')=='Verified']
        expected={'total_orders':len(order_rows),'revenue':sum((latest[o['id']]['amount'] for o in eligible),Decimal('0')),
                  'paid_orders':len(eligible),'today_orders':sum(bool(o['date']) and (o['date']+timedelta(hours=5,minutes=30)).date()==today for o in order_rows),
                  'refunded_orders':sum(o['refund_status']=='Completed' for o in order_rows),
                  'pending_payments':sum(o['status'] not in ('Cancelled','Refunded') and (latest.get(o['id'],{}).get('status') or o['payment_status'] or 'Pending') in ('Pending','Pending Verification','Pending COD') for o in order_rows),
                  'unverified_deliveries':sum(o['status']=='Delivered' and (latest.get(o['id'],{}).get('status') or o['payment_status'])!='Verified' for o in order_rows)}
        for key,status in [('pending_orders','Pending'),('delivered_orders','Delivered'),('cancelled_orders','Cancelled'),('transit_orders','Out For Delivery')]: expected[key]=sum(o['status']==status for o in order_rows)
        users=rows(self.db,'SELECT role,is_blocked FROM users')
        expected.update(total_users=len(users),customers=sum(u['role']=='customer' for u in users),staff_accounts=sum(u['role']=='staff' for u in users),blocked_users=sum(u['is_blocked']==1 for u in users))
        medicines=rows(self.db,'SELECT stock,low_stock_threshold,expiry_date FROM medicines')
        expected.update(medicines=len(medicines),available_stock=sum((m['stock'] or 0)>0 for m in medicines),
            low_stock=sum(m['stock'] is not None and 0<m['stock']<=(m['low_stock_threshold'] if m['low_stock_threshold'] is not None else 10) for m in medicines),
            out_of_stock=sum(m['stock'] is not None and m['stock']<=0 for m in medicines),unknown_stock=sum(m['stock'] is None for m in medicines),
            expired=sum(bool(m['expiry_date']) and m['expiry_date']<today for m in medicines),expiring=sum(bool(m['expiry_date']) and today<=m['expiry_date']<=today+timedelta(days=30) for m in medicines))
        prescriptions=rows(self.db,'SELECT status FROM prescription_requests');expected['prescriptions']=len(prescriptions)
        for key,status in [('pending_prescriptions','Pending Review'),('approved_prescriptions','Approved'),('rejected_prescriptions','Rejected')]: expected[key]=sum(p['status']==status for p in prescriptions)
        reviews=rows(self.db,'SELECT rating,status FROM reviews_feedback');ratings=[r['rating'] for r in reviews if r['rating'] is not None]
        expected.update(reviews=len(reviews),average_rating=Decimal(sum(ratings))/len(ratings) if ratings else None,submitted_feedback=sum(r['status']=='Submitted' for r in reviews))
        for key,table,where in [('active_subscriptions','medicine_subscriptions'," WHERE status='Active'"),('wishlist_entries','wishlist',''),('staff_profiles','staff','')]: expected[key]=rows(self.db,'SELECT COUNT(*) AS n FROM '+table+where)[0]['n']
        self.assertEqual(set(metrics),set(expected))
        for key,value in expected.items():
            with self.subTest(metric=key):
                if key=='average_rating' and value is not None:self.assertAlmostEqual(float(metrics[key]),float(value),places=4)
                else:self.assertEqual(metrics[key],value)
        print('Verified all',len(expected),'dashboard metrics; eligible revenue:',metrics['revenue'])

    def test_staff_fields_and_pagination(self):
        self.assertEqual(self.data['staff'], rows(self.db,'SELECT id,name,contact,address,email,age,gender,education FROM staff ORDER BY id DESC LIMIT 25'))
        self.assertLessEqual(len(self.data['orders']),25)
        second=load_dashboard(self.db,{'order_page':'2'})
        self.assertFalse(set(o['id'] for o in second['orders'])&set(o['id'] for o in self.data['orders']))

    def test_charts_and_report_totals(self):
        self.assertEqual(sum(r['value'] for r in self.data['status_rows']),self.data['metrics']['total_orders'])
        self.assertEqual(sum(r['value'] for r in self.data['payment_mix']),self.data['metrics']['revenue'])
        report=build_report(self.db,'sales',{})
        self.assertEqual(report['summary']['Net revenue'],self.data['metrics']['revenue'])
        self.assertEqual(report['summary']['Records'],self.data['metrics']['total_orders'])
        trend=self.data['revenue'];self.assertEqual(sum(trend['charts']['daily']['values']),trend['selected_period']['current'])

    def test_reports_all_formats_and_periods(self):
        from openpyxl import load_workbook
        today=now_ist().date()
        periods=[{}, {'from':str(today),'to':str(today)}, {'from':str(today-timedelta(days=6)),'to':str(today)},
                 {'from':'2026-06-01','to':'2026-08-31'}, {'from':'2000-01-01','to':'2000-01-02'}]
        for kind in ('sales','orders','inventory','prescriptions','users'):
            for period in periods if kind not in ('inventory','users') else [{}]:
                with self.subTest(kind=kind,period=period):
                    report=build_report(self.db,kind,period)
                    self.assertGreater(len(csv_bytes(report)),100)
                    self.assertTrue(pdf_bytes(report).startswith(b'%PDF-1.4'))
                    book=load_workbook(BytesIO(excel_bytes(report)))
                    self.assertEqual(book['Records'].max_row,max(1,len(report['rows']))+1)
                    if period.get('from')=='2000-01-01':self.assertEqual(report['rows'],[])

    def test_routes_and_access(self):
        import app as module
        app=module.app; app.config['TESTING']=True
        # Schema was independently inspected; read checks never run migration helpers.
        patches=[patch.object(module,name,return_value=None) for name in vars(module) if name.startswith('ensure_') and callable(getattr(module,name))]
        for p in patches:p.start()
        try:
            with patch.object(module,'get_db',return_value=self.db), app.test_client() as client:
                for role in (None,'customer','staff'):
                    with client.session_transaction() as session:
                        session.clear()
                        if role: session['user']={'id':-1,'role':role}
                    self.assertEqual(client.get('/owner_dashboard').status_code,302)
                    self.assertEqual(client.get('/owner/reports/sales').status_code,302)
                    self.assertEqual(client.get('/api/owner/revenue').status_code,401)
                with client.session_transaction() as session:session['user']={'id':-1,'role':'owner','name':'Owner'}
                for path in ['/owner_dashboard','/api/owner/revenue?filter=last_7_days','/owner/reports/sales',
                    '/owner/reports/orders?format=csv','/owner/reports/inventory?format=xlsx','/owner/reports/prescriptions?format=pdf',
                    '/owner/reports/users','/owner_dashboard?order_search=unlikely-no-match','/owner_dashboard?stock=out']:
                    with self.subTest(path=path):self.assertEqual(client.get(path).status_code,200)
                self.assertEqual(client.get('/owner_dashboard?order_page=-1').status_code,400)
                self.assertEqual(client.get('/owner/reports/sales?from=bad&to=bad').status_code,400)
                self.assertEqual(client.post('/remove_staff/1').status_code,400)
                with patch.object(module,'load_dashboard',side_effect=module.mysql.connector.Error('Simulated outage')):
                    self.assertEqual(client.get('/owner_dashboard').status_code,503)
        finally:
            for p in reversed(patches):p.stop()

if __name__=='__main__':unittest.main()
