"""Live, parameterized owner reports. No schema writes or stored report snapshots."""
import csv
from datetime import datetime, timedelta, time, timezone
from decimal import Decimal
from io import BytesIO, StringIO
import textwrap

from services.owner_dashboard import rows, PAYMENT_JOIN, ELIGIBLE, MONEY, IST, now_ist, display_date

REPORTS = {'sales': 'Sales Report', 'orders': 'Order Report', 'inventory': 'Inventory Report',
           'prescriptions': 'Prescription Report', 'users': 'User Report'}
MAX_ROWS = 10000


def date_filters(args):
    start, end = args.get('from', ''), args.get('to', '')
    if not start and not end:
        return None, None
    if not start or not end:
        raise ValueError('Choose both From and To dates.')
    try:
        start, end = (datetime.strptime(v, '%Y-%m-%d').date() for v in (start, end))
    except ValueError:
        raise ValueError('Dates must use YYYY-MM-DD.')
    if start > end or (end - start).days > 730:
        raise ValueError('Choose a date range of up to two years, with From before To.')
    return start, end


def build_report(db, kind, args):
    if kind not in REPORTS:
        raise ValueError('Unknown report type.')
    start, end = date_filters(args)
    if start and kind in ('inventory', 'users'):
        raise ValueError('This table has no registration/creation date; date filters are not supported.')
    where, params = ['1=1'], []
    status = args.get('status', '').strip()
    payment = args.get('payment_status', '').strip()
    if len(status) > 100 or len(payment) > 100:
        raise ValueError('Invalid status filter.')
    if status and kind not in ('sales', 'orders', 'prescriptions'):
        raise ValueError('Status filtering is not supported for this report.')
    if payment and kind not in ('sales', 'orders'):
        raise ValueError('Payment filtering is supported only for sales and orders.')
    date_column = 'o.date' if kind in ('sales', 'orders') else 'pr.created_at'
    if start:
        where += [f'{date_column} >= %s', f'{date_column} < %s']
        params += [datetime.combine(day, time.min, IST).astimezone(timezone.utc).replace(tzinfo=None)
                   for day in (start, end + timedelta(days=1))]
    if status:
        where.append(('pr.status' if kind == 'prescriptions' else 'o.status') + '=%s')
        params.append(status)
    if payment:
        where.append('COALESCE(p.status,o.payment_status)=%s'); params.append(payment)
    if kind in ('sales', 'orders'):
        common = f'''FROM orders o LEFT JOIN users u ON u.id=o.user_id {PAYMENT_JOIN}
            LEFT JOIN (SELECT order_id, SUM(quantity) AS items,
                SUM(CAST(price AS DECIMAL(14,2))*quantity) AS subtotal
                FROM order_items GROUP BY order_id) oi ON oi.order_id=o.id'''
        if kind == 'sales':
            fields = '''o.date AS date,o.id AS order_id,u.name AS customer,COALESCE(oi.items,0) AS items,
                oi.subtotal AS item_subtotal,o.delivery_charge AS delivery_fee,
                CAST(o.total AS DECIMAL(14,2)) AS final_total,COALESCE(p.method,o.payment_method) AS payment_method,
                COALESCE(p.status,o.payment_status) AS payment_status,o.status AS order_status'''
            columns = [('date','Order Date'),('order_id','Order ID'),('customer','Customer'),('items','Items'),
                       ('item_subtotal','Item Subtotal'),('delivery_fee','Delivery Fee'),('final_total','Final Total'),
                       ('payment_method','Payment Method'),('payment_status','Payment Status'),('order_status','Order Status'),('net_revenue','Net Revenue')]
        else:
            fields = '''o.id AS order_id,u.name AS customer,o.date AS date,
                CAST(o.total AS DECIMAL(14,2)) AS final_total,
                COALESCE(p.status,o.payment_status) AS payment_status,o.status AS order_status,
                CASE WHEN NULLIF(o.prescription,'') IS NOT NULL THEN 'Attached' ELSE 'Not attached' END AS prescription,
                o.delivery_address,o.delivery_area,o.delivery_status'''
            columns = [('order_id','Order ID'),('customer','Customer'),('date','Order Date'),('final_total','Total'),
                       ('payment_status','Payment Status'),('order_status','Order Status'),('prescription','Prescription'),
                       ('delivery_address','Delivery Address'),('delivery_area','Delivery Area'),('delivery_status','Delivery Status')]
        sql = f'''SELECT {fields}, CASE WHEN {ELIGIBLE} THEN {MONEY} ELSE 0 END AS net_revenue
            {common} WHERE {' AND '.join(where)} ORDER BY o.date DESC,o.id DESC'''
    elif kind == 'inventory':
        columns = [('id','Medicine ID'),('name','Medicine'),('category','Category'),('stock','Current Stock'),
                   ('price','Selling Price'),('low_stock_threshold','Low Stock Threshold'),('expiry_date','Expiry Date'),('stock_status','Stock Status')]
        sql = '''SELECT id,name,category,stock,CAST(price AS DECIMAL(14,2)) AS price,low_stock_threshold,expiry_date,
            CASE WHEN stock IS NULL THEN 'Unknown' WHEN stock <= 0 THEN 'Out of stock'
                 WHEN stock <= COALESCE(low_stock_threshold,10) THEN 'Low stock' ELSE 'Available' END AS stock_status
            FROM medicines ORDER BY id'''
    elif kind == 'prescriptions':
        columns = [('id','Request ID'),('customer','Customer'),('created_at','Requested Date'),('prescription_image','Prescription File'),
                   ('status','Status'),('reviewed_at','Review Date'),('staff_note','Review Note'),('created_order_id','Created Order ID')]
        sql = f'''SELECT pr.id,u.name AS customer,pr.created_at,pr.prescription_image,pr.status,pr.reviewed_at,
            pr.staff_note,pr.created_order_id FROM prescription_requests pr LEFT JOIN users u ON u.id=pr.user_id
            WHERE {' AND '.join(where)} ORDER BY pr.created_at DESC,pr.id DESC'''
    else:
        columns = [('id','User ID'),('name','Name'),('email','Email'),('phone','Phone'),('role','Role'),('account_status','Account Status')]
        sql = "SELECT id,name,email,phone,role,CASE WHEN is_blocked=1 THEN 'Blocked' ELSE 'Active' END AS account_status FROM users ORDER BY id"
    data = rows(db, sql + ' LIMIT %s', [*params, MAX_ROWS + 1])
    if len(data) > MAX_ROWS:
        raise ValueError('Report exceeds 10,000 records. Narrow the date range or status filter.')
    summary = {'Records': len(data)}
    if kind == 'sales':
        for key, label in [('item_subtotal','Gross item sales'),('delivery_fee','Delivery charges'),('final_total','Order value'),('net_revenue','Net revenue')]:
            summary[label] = sum((Decimal(str(r[key] or 0)) for r in data), Decimal('0'))
    notes = {
        'sales': 'Order-date cohort. Net revenue uses the latest verified payment and excludes cancelled/refunded orders and refunds in progress. Item subtotal is calculated from saved order lines. Discounts are not stored separately and are not fabricated; order total is authoritative. Currency: INR.',
        'orders': 'Prescription indicates a saved attachment; historical prescription requirement is not stored on order lines. Currency: INR.',
        'inventory': 'Current inventory snapshot. Threshold uses the stored low_stock_threshold (legacy fallback: 10). Currency: INR.',
        'prescriptions': 'Reviewer identity is not stored on prescription_requests and is therefore omitted.',
        'users': 'Registration date is not stored. Account status is based on is_blocked. Authentication secrets are excluded.',
    }
    return dict(title=REPORTS[kind], kind=kind, columns=columns, rows=data, summary=summary,
                generated_at=now_ist(), period=f'{start} to {end}' if start else 'All records / current snapshot',
                filters={k: args.get(k,'') for k in ('from','to','status','payment_status')}, note=notes[kind])


def cell(value):
    if value is None:
        return '—'
    if isinstance(value, (datetime,)) or hasattr(value, 'strftime'):
        return display_date(value)
    if isinstance(value, Decimal):
        return f'{value:.2f}'
    return str(value)


def spreadsheet_cell(value):
    if isinstance(value, (int, float, Decimal)):
        return value
    text = cell(value)
    # Prevent spreadsheet applications interpreting user text as formulas.
    return "'" + text if text.lstrip().startswith(('=', '+', '-', '@', '\t', '\r')) else text


def csv_bytes(report):
    stream = StringIO(newline='')
    writer = csv.writer(stream)
    writer.writerow(['YuvrajMedical', report['title']])
    writer.writerow(['Generated (IST)', display_date(report['generated_at'])])
    writer.writerow(['Period', report['period']])
    writer.writerow(['Filters', spreadsheet_cell(str(report['filters']))])
    writer.writerow([report['note']]); writer.writerow([])
    writer.writerow([label for _, label in report['columns']])
    for record in report['rows']:
        writer.writerow([spreadsheet_cell(record[key]) for key, _ in report['columns']])
    if not report['rows']: writer.writerow(['No matching records'])
    writer.writerow([])
    for key, value in report['summary'].items(): writer.writerow([key, value])
    return stream.getvalue().encode('utf-8-sig')


def excel_bytes(report):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    sheet = workbook.active; sheet.title = 'Records'
    sheet.append([label for _, label in report['columns']])
    for record in report['rows']:
        sheet.append([spreadsheet_cell(record[key]) for key, _ in report['columns']])
    if not report['rows']: sheet.append(['No matching records'])
    sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
    for c in sheet[1]: c.font = Font(color='FFFFFF', bold=True); c.fill = PatternFill('solid', fgColor='087F8C')
    for i, _ in enumerate(report['columns'], 1): sheet.column_dimensions[get_column_letter(i)].width = 23
    for row in sheet.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(c.value, (float, Decimal)): c.number_format = '#,##0.00'
    summary = workbook.create_sheet('Summary')
    for record in [('YuvrajMedical', report['title']), ('Generated (IST)', display_date(report['generated_at'])),
                   ('Period', report['period']), ('Filters', spreadsheet_cell(str(report['filters']))), ('Definition', report['note']), *report['summary'].items()]: summary.append(record)
    summary.column_dimensions['A'].width = 25; summary.column_dimensions['B'].width = 100
    output = BytesIO(); workbook.save(output); return output.getvalue()


def pdf_bytes(report):
    """Dependency-free paginated landscape PDF with repeated column headings."""
    columns = report['columns']
    width = max(10, 142 // len(columns) - 1)
    def ascii_text(value):
        return cell(value).replace('₹','INR ').replace('—','-').encode('ascii','replace').decode('ascii')
    def wrapped_cells(values):
        wrapped = [textwrap.wrap(ascii_text(v), width) or [''] for v in values]
        return [' '.join((part[i] if i < len(part) else '').ljust(width) for part in wrapped)
                for i in range(max(map(len,wrapped)))]
    headings = wrapped_cells([label for _, label in columns])
    pages, current = [], []
    def add(line):
        nonlocal current
        if len(current) >= 37:
            pages.append(current); current = []
        current.append(line)
    for record in report['rows']:
        for line in wrapped_cells([record[k] for k, _ in columns]): add(line)
        add('')
    if not report['rows']: add('No matching records.')
    for key, value in report['summary'].items(): add(f'{key}: {ascii_text(value)}')
    if current: pages.append(current)
    objects = [b'', b'', b'<< /Type /Font /Subtype /Type1 /BaseFont /Courier >>']
    page_ids = []
    for number, lines in enumerate(pages, 1):
        header = [f'YuvrajMedical | {report["title"]} | Page {number}/{len(pages)}',
                  f'Generated (IST): {display_date(report["generated_at"])} | {report["period"]}',
                  f'Filters: {report["filters"]}', *textwrap.wrap(report['note'], 140), '', *headings, '-' * min(142,(width+1)*len(columns))]
        commands = ['BT', '/F1 8 Tf', '28 566 Td', '10 TL']
        for line in header + lines:
            text = ascii_text(line).replace('\\','\\\\').replace('(','\\(').replace(')','\\)')
            commands += [f'({text}) Tj', 'T*']
        commands.append('ET'); stream = '\n'.join(commands).encode('ascii')
        page_id = len(objects) + 1; content_id = page_id + 1; page_ids.append(page_id)
        objects.append(f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 842 595] /Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>'.encode())
        objects.append(b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream')
    objects[0] = b'<< /Type /Catalog /Pages 2 0 R >>'
    objects[1] = f'<< /Type /Pages /Kids [{" ".join(str(i)+" 0 R" for i in page_ids)}] /Count {len(page_ids)} >>'.encode()
    output=BytesIO(); output.write(b'%PDF-1.4\n'); offsets=[0]
    for i,obj in enumerate(objects,1):
        offsets.append(output.tell()); output.write(f'{i} 0 obj\n'.encode()+obj+b'\nendobj\n')
    xref=output.tell(); output.write(f'xref\n0 {len(offsets)}\n0000000000 65535 f \n'.encode())
    for offset in offsets[1:]: output.write(f'{offset:010d} 00000 n \n'.encode())
    output.write(f'trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode())
    return output.getvalue()
